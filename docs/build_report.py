from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "AI_Phishing_Detector_Report.docx"
METRICS = ROOT / "model" / "metrics.txt"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(85, 85, 85)
RISK_RED = RGBColor(155, 28, 28)
GREEN = RGBColor(22, 101, 52)
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="DADCE0", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_table(table, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for row_idx, row in enumerate(table.rows):
        if row_idx == 0 and header:
            set_repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx == 0 and header:
                set_cell_fill(cell, LIGHT_GRAY)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = DARK_BLUE


def set_column_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_field(paragraph, field):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr)
    run._r.append(fld_char_end)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_footer(doc):
    footer = doc.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("AI Phishing Detector Report | Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    add_field(paragraph, "PAGE")


def add_title(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AI Browser Extension for Phishing Detection")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = DARK_BLUE

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("Final Project Report")
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = MUTED

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "23K-0714 S M Attique Ur Rehman | 23K-0574 Anfas Ali | 23K-0511 Asher Ahmed"
    )
    meta_run.font.size = Pt(10)
    meta_run.font.color.rgb = MUTED

    doc.add_paragraph()


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_fill(cell, CALLOUT)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    set_table_borders(table, color="DADCE0")
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    p.add_run(f" {body}")
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(item)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(item)


def metrics_summary():
    if not METRICS.exists():
        return "Metrics file was not found. Run python -m backend.train_model to regenerate metrics."
    text = METRICS.read_text(encoding="utf-8", errors="ignore")
    lines = []
    capture = False
    for line in text.splitlines():
        if line.startswith("Best model:"):
            lines.append(line)
        if line.startswith("Model: Random Forest"):
            capture = True
        if capture:
            lines.append(line)
            if line.strip().startswith("phishing"):
                break
    return "\n".join(lines[:16])


def add_architecture_table(doc):
    table = doc.add_table(rows=1, cols=3)
    headers = ["Component", "Main Files", "Purpose"]
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
    rows = [
        ("Browser Extension", "extension/manifest.json, background.js, popup.js, content.js", "Captures visited URLs, displays scan status, and shows phishing warnings."),
        ("Flask Backend", "backend/app.py", "Provides /predict and /mark-legitimate APIs for scanning and user feedback."),
        ("Feature Extraction", "backend/features.py", "Converts URLs into numeric ML features such as entropy, suspicious keywords, IP address flags, and query structure."),
        ("Model Training", "backend/train_model.py", "Trains and compares ML models, then saves the best classifier as phishing_model.pkl."),
        ("Datasets", "dataset/*.csv, real_phishdataset_balanced.xlsx", "Stores real phishing data, trusted legitimate domains, and user feedback."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    set_column_widths(table, [1.45, 2.25, 2.8])
    format_table(table)


def add_results_table(doc):
    table = doc.add_table(rows=1, cols=4)
    headers = ["Evaluation Item", "Result", "Interpretation", "Status"]
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
    rows = [
        ("Selected model", "Random Forest", "Best model according to phishing recall and F1-score.", "Passed"),
        ("Accuracy", "Approx. 94.5%", "Correctly classifies most URLs in the held-out split.", "Passed"),
        ("Phishing recall", "Approx. 94.9%", "Prioritizes catching phishing URLs, reducing missed attacks.", "Passed"),
        ("Backend tests", "8 passed", "API, trusted domains, and feedback endpoint verified.", "Passed"),
        ("Extension checks", "JS syntax passed", "Popup, background, and content scripts parse correctly.", "Passed"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    set_column_widths(table, [1.5, 1.3, 2.8, 0.9])
    format_table(table)


def build_report():
    doc = Document()
    configure_document(doc)
    add_footer(doc)
    add_title(doc)

    add_callout(
        doc,
        "Executive summary.",
        "This project implements a Chrome/Edge browser extension that detects phishing URLs in real time using a Flask backend and a machine learning classifier. It warns users before they enter sensitive information and includes user feedback to reduce false positives.",
    )

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "Phishing websites imitate trusted services to steal credentials, card details, and personal information. The goal of this project is to demonstrate how artificial intelligence can support safer browsing by classifying URLs as legitimate or suspicious while the user browses the web."
    )
    doc.add_paragraph(
        "The solution combines machine learning, URL feature engineering, a Flask prediction API, and a Manifest V3 browser extension. The extension performs real-time checks, displays a popup status, and shows a full-page warning when a phishing risk is detected."
    )

    doc.add_heading("2. Objectives", level=1)
    add_bullets(
        doc,
        [
            "Detect phishing websites in real time from the browser.",
            "Train and evaluate machine learning models for URL classification.",
            "Warn users clearly before they enter sensitive information.",
            "Reduce false positives through trusted domains and user feedback.",
            "Demonstrate a practical AI application in cybersecurity.",
        ],
    )

    doc.add_heading("3. System Architecture", level=1)
    doc.add_paragraph(
        "The system is divided into five main parts. The extension captures the current tab URL, the backend predicts risk, the model performs classification, and feedback files allow user corrections to improve future scans."
    )
    add_architecture_table(doc)

    doc.add_heading("4. Dataset and Feature Engineering", level=1)
    doc.add_paragraph(
        "The project uses the public ESDAUNG/PhishDataset balanced dataset, based on PhishTank phishing URLs and legitimate URL sources. It is normalized into url,label format. The project also includes trusted legitimate domains and user feedback data to reduce false positives for well-known services."
    )
    doc.add_heading("Key URL Features", level=2)
    add_bullets(
        doc,
        [
            "Length-based features: URL length, domain length, path length, query length, and longest token length.",
            "Character-based features: dot count, hyphen count, digit count, special character ratio, and entropy.",
            "Security indicators: HTTPS usage, IP address in hostname, suspicious TLD, URL shortener, and encoded characters.",
            "Phishing patterns: suspicious keywords, brand keywords in subdomains, executable file extensions, and redirect-like paths.",
        ],
    )

    doc.add_heading("5. Machine Learning Model", level=1)
    doc.add_paragraph(
        "The training script evaluates Logistic Regression, Decision Tree, Random Forest, and Naive Bayes. The selected model is chosen using phishing recall first and F1-score second because missing a phishing website is more dangerous than marking an extra site as suspicious."
    )
    add_callout(doc, "Current model result.", metrics_summary().replace("\n", " | "))
    add_results_table(doc)

    doc.add_heading("6. Backend API", level=1)
    doc.add_paragraph(
        "The Flask backend exposes two main endpoints. The /predict endpoint receives a URL, checks trusted domains, extracts features, loads the saved model, and returns a prediction with risk score. The /mark-legitimate endpoint lets the user correct false positives."
    )
    add_numbered(
        doc,
        [
            "Extension sends the current URL to /predict.",
            "Backend checks trusted legitimate domains before calling the model.",
            "If not trusted, features are extracted and passed to the trained model.",
            "Backend returns prediction, risk score, URL, and source where applicable.",
            "If the user marks a false positive as legitimate, the URL and domain are saved for future scans.",
        ],
    )

    doc.add_heading("7. Browser Extension", level=1)
    doc.add_paragraph(
        "The extension uses Chrome Manifest V3. Its background script monitors tab updates, calls the Flask backend, stores results, and notifies content scripts. The popup displays the scan result, risk score, and actions. The content script renders a full-page warning overlay for suspicious websites."
    )
    doc.add_heading("User Interface Behavior", level=2)
    add_bullets(
        doc,
        [
            "Safe sites show a green safe status and low risk score.",
            "Suspicious sites show a red danger status and warning text.",
            "The warning overlay includes Go Back, Mark as Legit, and Continue Anyway.",
            "Mark as Legit records user feedback and immediately trusts the corrected domain.",
            "Rescan Current Page sends the current URL to the backend again without retraining the model.",
        ],
    )

    doc.add_heading("8. Testing", level=1)
    doc.add_paragraph(
        "Testing covered feature extraction, backend endpoints, trusted-domain handling, user feedback, JavaScript syntax, and manual browser behavior."
    )
    add_bullets(
        doc,
        [
            "Feature tests verify URL normalization and suspicious signal extraction.",
            "Backend tests verify /health, /predict, trusted domains, and /mark-legitimate.",
            "JavaScript syntax checks verify background.js, popup.js, and content.js.",
            "Manual tests verify safe websites, suspicious phishing-like URLs, warning overlay actions, and user feedback.",
        ],
    )

    doc.add_heading("9. Limitations", level=1)
    add_bullets(
        doc,
        [
            "The model uses URL-based features only and does not inspect full webpage HTML or screenshots.",
            "The extension depends on a locally running Flask backend.",
            "Trusted-domain overrides reduce false positives but must be maintained carefully to avoid over-trusting broad domains.",
            "The model is suitable for academic demonstration and should not be treated as production security software without further validation.",
        ],
    )

    doc.add_heading("10. Future Improvements", level=1)
    add_bullets(
        doc,
        [
            "Add webpage content features such as forms, password fields, hidden inputs, and suspicious scripts.",
            "Use domain age, certificate information, WHOIS data, and reputation services.",
            "Add automated retraining after enough user feedback is collected.",
            "Package the model inside the extension using ONNX or TensorFlow.js to remove the Flask dependency.",
            "Add a dashboard showing scan history, blocked sites, and false-positive corrections.",
        ],
    )

    doc.add_heading("11. Conclusion", level=1)
    doc.add_paragraph(
        "The project demonstrates a complete AI-assisted phishing detection workflow. It combines a browser extension, Flask API, machine learning model, real phishing dataset, trusted-domain handling, user feedback, and clear warning actions. The result is an interactive cybersecurity tool that can identify risky URLs and help users avoid phishing attacks."
    )

    doc.add_heading("References", level=1)
    add_bullets(
        doc,
        [
            "ESDAUNG/PhishDataset public dataset: https://github.com/ESDAUNG/PhishDataset",
            "PhishTank phishing URL source referenced by the dataset: https://phishtank.org",
            "Scikit-learn machine learning library: https://scikit-learn.org",
            "Chrome Extensions Manifest V3 documentation: https://developer.chrome.com/docs/extensions",
            "Flask web framework documentation: https://flask.palletsprojects.com",
        ],
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
